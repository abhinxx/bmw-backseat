# E_01: dump full raw structure of both sheets to a text file for inspection.
import openpyxl, json, sys

SRC = "BMW-Data/F70_G60_G68_G70_customer_studies.xlsx"
wb = openpyxl.load_workbook(SRC, data_only=True)

out = open("research/data/E_raw_dump.txt", "w", encoding="utf-8")
for ws in wb.worksheets:
    out.write(f"=== SHEET {ws.title} dims={ws.dimensions} merged={list(ws.merged_cells.ranges)}\n")
    for row in ws.iter_rows():
        vals = []
        for c in row:
            v = c.value
            if v is None:
                continue
            vals.append(f"{c.coordinate}={v!r}")
        if vals:
            out.write(" | ".join(vals) + "\n")
out.close()

# also check formulas (non-data_only) to know if values are computed
wb2 = openpyxl.load_workbook(SRC, data_only=False)
fo = open("research/data/E_formulas.txt", "w", encoding="utf-8")
n = 0
for ws in wb2.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                fo.write(f"{ws.title}!{c.coordinate}: {c.value}\n")
                n += 1
fo.write(f"total formulas: {n}\n")
fo.close()
print("done; formulas:", n)
